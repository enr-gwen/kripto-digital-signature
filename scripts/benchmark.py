"""Benchmark Hari 5: waktu sign/verify, ukuran signature & kunci, uji
keamanan (tamper, kunci salah, QR palsu). Menulis hasil ke XLSX + grafik.

Jalankan dari folder utama proyek:
    python3 scripts/benchmark.py
"""
import statistics
import sys
import time
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import pandas as pd  # noqa: E402
from reportlab.pdfgen import canvas as _canvas  # noqa: E402

from app import keys, qr, service, signer  # noqa: E402
from app.pdf_stamp import stamp_qr  # noqa: E402

N_PERCOBAAN = 30
OUT_DIR = ROOT / "hasil_uji"
OUT_DIR.mkdir(exist_ok=True)


def buat_pdf_contoh(nama_file: str, ukuran_kb_kira: int = 5) -> bytes:
    import io
    buf = io.BytesIO()
    c = _canvas.Canvas(buf, pagesize=(595.27, 841.89))
    baris = max(1, ukuran_kb_kira // 1)  # kasar, cukup untuk variasi ukuran
    for i in range(baris):
        c.drawString(72, 800 - (i % 40) * 18,
                     f"Baris pengujian ke-{i} untuk {nama_file} " * 3)
        if (i + 1) % 40 == 0:
            c.showPage()
    c.save()
    return buf.getvalue()


def ukur_waktu(fn, n=N_PERCOBAAN):
    waktu = []
    for _ in range(n):
        mulai = time.perf_counter()
        fn()
        waktu.append((time.perf_counter() - mulai) * 1000)  # ms
    return waktu


def ringkas(label, waktu_ms):
    return {
        "Operasi": label,
        "Jumlah Percobaan": len(waktu_ms),
        "Rata-rata (ms)": round(statistics.mean(waktu_ms), 4),
        "Median (ms)": round(statistics.median(waktu_ms), 4),
        "Std Dev (ms)": round(statistics.stdev(waktu_ms), 4),
        "Min (ms)": round(min(waktu_ms), 4),
        "Max (ms)": round(max(waktu_ms), 4),
    }


def main():
    print(f"Menjalankan benchmark ({N_PERCOBAAN}x percobaan)...")
    private_key = keys.generate_keypair()
    public_key = private_key.public_key()
    pdf_uji = buat_pdf_contoh("dokumen_uji", ukuran_kb_kira=5)
    digest = signer.hash_bytes(pdf_uji)
    signature_contoh = signer.sign_digest(private_key, digest)

    # 1) Waktu sign & verify (level fungsi kripto, tanpa I/O file/DB)
    waktu_sign = ukur_waktu(lambda: signer.sign_digest(private_key, digest))
    waktu_verify = ukur_waktu(
        lambda: signer.verify_digest(public_key, digest, signature_contoh))

    # 2) Waktu alur lengkap (sign_document & verify_document, termasuk QR+DB)
    db_path = Path(tempfile.mktemp(suffix=".db"))

    def alur_sign():
        service.sign_document(
            pdf_uji, nama="Budi Santoso", jabatan="Ketua Panitia",
            institusi="Universitas Siliwangi", private_key=private_key,
            db_path=db_path)

    waktu_alur_sign = ukur_waktu(alur_sign)
    pdf_tertanda, doc_id = service.sign_document(
        pdf_uji, nama="Budi Santoso", jabatan="Ketua Panitia",
        institusi="Universitas Siliwangi", private_key=private_key,
        db_path=db_path)
    waktu_alur_verify = ukur_waktu(
        lambda: service.verify_document(
            pdf_tertanda, public_key=public_key, db_path=db_path))

    df_waktu = pd.DataFrame([
        ringkas("Sign (fungsi kripto saja)", waktu_sign),
        ringkas("Verify (fungsi kripto saja)", waktu_verify),
        ringkas("Sign dokumen (alur lengkap: QR+stempel+DB)", waktu_alur_sign),
        ringkas("Verify dokumen (alur lengkap: baca QR+DB)", waktu_alur_verify),
    ])

    # 3) Ukuran signature & kunci
    pub_pem = keys.public_key_to_pem(private_key)
    priv_pem = keys.private_key_to_pem(private_key, b"contoh-benchmark")
    df_ukuran = pd.DataFrame([
        {"Item": "Signature ECDSA P-256 (DER)", "Ukuran (byte)": len(signature_contoh)},
        {"Item": "Public key (PEM)", "Ukuran (byte)": len(pub_pem)},
        {"Item": "Private key terenkripsi (PEM)", "Ukuran (byte)": len(priv_pem)},
        {"Item": "PDF sebelum stempel QR", "Ukuran (byte)": len(pdf_uji)},
        {"Item": "PDF sesudah stempel QR", "Ukuran (byte)": len(pdf_tertanda)},
        {"Item": "Selisih akibat stempel QR", "Ukuran (byte)":
            len(pdf_tertanda) - len(pdf_uji)},
    ])

    # 4) Uji keamanan/fungsional (bukti kualitatif, bukan waktu)
    hasil_valid = service.verify_document(
        pdf_tertanda, public_key=public_key, db_path=db_path)

    rusak = bytearray(pdf_tertanda)
    rusak[len(rusak) // 2] ^= 0x01
    hasil_tamper = service.verify_document(
        bytes(rusak), public_key=public_key, db_path=db_path, doc_id=doc_id)

    kunci_lain = keys.generate_keypair()
    hasil_kunci_salah = service.verify_document(
        pdf_tertanda, public_key=kunci_lain.public_key(), db_path=db_path)

    payload_palsu = qr.build_payload("ffffffffffffffff", "Penipu", "Direktur",
                                     "2026-01-01", "Institusi Palsu",
                                     "http://localhost:8501")
    img_palsu = qr.make_qr_image(qr.payload_to_text(payload_palsu))
    pdf_qr_palsu = stamp_qr(pdf_uji, img_palsu, ["Palsu"])
    hasil_qr_palsu = service.verify_document(
        pdf_qr_palsu, public_key=public_key, db_path=db_path)

    df_keamanan = pd.DataFrame([
        {"Skenario Uji": "Dokumen asli, kunci benar",
         "Hasil Diharapkan": "VALID", "Hasil Aktual": hasil_valid.status,
         "Sesuai?": hasil_valid.status == "VALID"},
        {"Skenario Uji": "Dokumen diubah 1 byte (tamper)",
         "Hasil Diharapkan": "DOKUMEN_DIUBAH", "Hasil Aktual": hasil_tamper.status,
         "Sesuai?": hasil_tamper.status == "DOKUMEN_DIUBAH"},
        {"Skenario Uji": "Kunci publik salah",
         "Hasil Diharapkan": "KUNCI_TIDAK_COCOK",
         "Hasil Aktual": hasil_kunci_salah.status,
         "Sesuai?": hasil_kunci_salah.status == "KUNCI_TIDAK_COCOK"},
        {"Skenario Uji": "QR palsu (doc_id tidak dikenal)",
         "Hasil Diharapkan": "DOC_ID_TIDAK_DIKENAL",
         "Hasil Aktual": hasil_qr_palsu.status,
         "Sesuai?": hasil_qr_palsu.status == "DOC_ID_TIDAK_DIKENAL"},
    ])

    # 5) Simpan mentah waktu (untuk lampiran/analisis lanjutan)
    df_mentah = pd.DataFrame({
        "Percobaan": range(1, N_PERCOBAAN + 1),
        "Sign (ms)": waktu_sign,
        "Verify (ms)": waktu_verify,
        "Sign Dokumen (ms)": waktu_alur_sign,
        "Verify Dokumen (ms)": waktu_alur_verify,
    })

    out_path = OUT_DIR / "hasil_pengujian.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df_waktu.to_excel(writer, sheet_name="Ringkasan Waktu", index=False)
        df_ukuran.to_excel(writer, sheet_name="Ukuran", index=False)
        df_keamanan.to_excel(writer, sheet_name="Uji Keamanan", index=False)
        df_mentah.to_excel(writer, sheet_name="Data Mentah", index=False)

    _tambah_grafik(out_path, df_mentah)

    print(f"\nSelesai. Hasil tersimpan di: {out_path}\n")
    print(df_waktu.to_string(index=False))
    print()
    print(df_ukuran.to_string(index=False))
    print()
    print(df_keamanan.to_string(index=False))
    if not df_keamanan["Sesuai?"].all():
        print("\n⚠️  ADA UJI KEAMANAN YANG TIDAK SESUAI HARAPAN! Cek tabel di atas.")
        sys.exit(1)


def _tambah_grafik(xlsx_path, df_mentah):
    """Tambah sheet Grafik berisi chart batang (openpyxl), dibuat dari sheet
    'Data Mentah' yang sudah ditulis pandas."""
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.load_workbook(xlsx_path)
    ws_data = wb["Data Mentah"]
    ws_chart = wb.create_sheet("Grafik")

    chart = BarChart()
    chart.title = "Waktu Sign vs Verify per Percobaan (fungsi kripto)"
    chart.x_axis.title = "Percobaan ke-"
    chart.y_axis.title = "Waktu (ms)"
    n = len(df_mentah) + 1  # +1 header

    data_sign = Reference(ws_data, min_col=2, max_col=2, min_row=1, max_row=n)
    data_verify = Reference(ws_data, min_col=3, max_col=3, min_row=1, max_row=n)
    kategori = Reference(ws_data, min_col=1, min_row=2, max_row=n)
    chart.add_data(data_sign, titles_from_data=True)
    chart.add_data(data_verify, titles_from_data=True)
    chart.set_categories(kategori)
    chart.width, chart.height = 24, 12
    ws_chart.add_chart(chart, "A1")
    wb.save(xlsx_path)


if __name__ == "__main__":
    main()
